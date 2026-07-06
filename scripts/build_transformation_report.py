from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "BusLoop_Project_Transformation_Report.docx"
LOGO = ROOT / "busnow-user-app" / "public" / "busloop-mark.png"

NAVY = RGBColor(11, 37, 69)
RED = RGBColor(239, 61, 65)
GREEN = RGBColor(18, 183, 106)
BLUE = RGBColor(37, 99, 235)
MUTED = RGBColor(98, 111, 134)
HEADER = "E8EEF5"
SUCCESS = "EAF8F1"
WARNING = "FFF7E6"
DANGER = "FDECEC"
LIGHT = "F7F9FC"


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_text(cell, text, bold=False, color=None, size=9.3):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color


def set_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def table(doc, headers, rows, widths=None, fill=HEADER):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_shading(t.rows[0].cells[i], fill)
        set_text(t.rows[0].cells[i], h, bold=True, color=NAVY)
    for row in rows:
        cells = t.add_row().cells
        for i, item in enumerate(row):
            set_text(cells[i], item)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if widths:
        set_widths(t, widths)
    doc.add_paragraph()
    return t


def para(doc, text="", bold=False, color=None, size=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "Calibri"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        r.font.color.rgb = NAVY if level == 1 else BLUE
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def callout(doc, title, body, fill=SUCCESS):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    set_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = NAVY
    r.font.size = Pt(11)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(body)
    set_widths(t, [6.25])
    doc.add_paragraph()


def setup(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in [
        ("Heading 1", 16, NAVY),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11.5, NAVY),
    ]:
        s = doc.styles[style_name]
        s.font.name = "Calibri"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = color

    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("BusLoop Transformation Report")
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = MUTED
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("BusLoop | Project Recovery, Stabilization, and Deployment Readiness")
    fr.font.size = Pt(8.5)
    fr.font.color.rgb = MUTED


def build():
    doc = Document()
    setup(doc)

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.4))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("BusLoop Project Transformation Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Initial condition, issues found, work completed, and current deployment-ready state")
    sr.font.size = Pt(11.5)
    sr.font.color.rgb = MUTED

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("Prepared for client handover | July 6, 2026")
    mr.font.size = Pt(10)
    mr.font.color.rgb = MUTED

    callout(
        doc,
        "Short Version",
        "The project was received in an incomplete, inconsistent, and demo-risky state. It had broken payment and ticket flows, unreliable verification, missing QR use cases, inconsistent UI across passenger and staff apps, weak role handling, and several production-readiness gaps. It has now been stabilized, rebranded as BusLoop, deployed, tested end to end, and documented for client demo and real-time pilot use.",
        fill=SUCCESS,
    )

    table(
        doc,
        ["Current Asset", "Link / Status"],
        [
            ["Passenger App", "https://busloop-user-app.vercel.app"],
            ["Staff App", "https://busloop-staff-app.vercel.app"],
            ["GitHub Repository", "https://github.com/imadi005/BusLoop"],
            ["Client E2E Report", "BusLoop_Client_Project_Report.docx"],
            ["Current Status", "Ready for deployment, client demo, and real-time pilot use"],
        ],
        widths=[1.8, 4.45],
    )

    doc.add_page_break()

    heading(doc, "1. Condition When the Project Was Received")
    para(doc, "The project was not in a production-ready or client-demo-ready condition. Multiple major workflows were either incomplete, broken, inconsistent, or only partially implemented. The app looked like it had been built in separate pieces without enough end-to-end testing.")
    bullets(
        doc,
        [
            "Payment flow was not properly integrated into the full user workflow.",
            "Ticket generation after payment needed to be connected and verified.",
            "Ticket verification modal and checker verification flow were not working reliably.",
            "Passenger and staff apps had very different visual styles, making the product feel inconsistent.",
            "Important user cases were missing, including passenger scanning a bus QR to buy a ticket.",
            "Staff role handling was unclear; checker role visibility and privileges were confusing.",
            "Camera QR scanner behavior was broken or showing poor fallback screens.",
            "Operator, checker, and driver pages had unnecessary noise, inconsistent layouts, and confusing states.",
            "Some calculations and status indicators were logically wrong, such as percentage/rate values with zero data.",
            "Ticket validity, expiry, boarding time, and refund/cancellation rules were missing or unclear.",
            "Profile and review areas on the passenger side were incomplete and not easy to use.",
            "Live bus and QR test cases were not fully prepared for realistic client demo.",
            "Branding was still old and inconsistent before the final BusLoop rebrand.",
        ],
    )

    heading(doc, "2. Main Problems Found")
    table(
        doc,
        ["Area", "Problem Found", "Impact"],
        [
            ["Payment", "Razorpay was not fully connected to ticket generation and workflow testing.", "Client could not confidently demo paid ticket booking."],
            ["Ticketing", "Ticket modal lag, scroll issues, missing validity rules, and QR display issues.", "Passenger experience felt unfinished."],
            ["Verification", "Checker scan and manual verification were unreliable.", "Boarding use case was not dependable."],
            ["Roles", "Checker role and permissions were unclear; staff role routing needed correction.", "Staff app could show wrong role/state."],
            ["Passenger QR", "Bus QR scan purchase case had been removed/missing.", "A key real-world bus onboarding case was absent."],
            ["Camera", "Camera fallback showed generic/bad UI and mobile browser cases failed.", "QR scanning felt broken during phone testing."],
            ["Operator UI", "Dashboard, analytics, buses, and routes looked inconsistent and cluttered.", "Operator side did not feel professional."],
            ["Checker UI", "Unnecessary route/noise and incorrect rate logic.", "Verification workflow was harder to understand."],
            ["Driver UI", "Driver states and GPS workflow needed cleanup and testing.", "Live tracking was not demo-safe."],
            ["Reviews", "Redundant and confusing review areas.", "Passenger review UX was unclear."],
            ["Brand", "Old BusNOW branding remained in app and deployment URLs.", "Client-facing identity was not ready."],
            ["Deployment", "Vercel URLs and app names still used old BusNOW naming.", "Final handover looked incomplete."],
        ],
        widths=[1.25, 3.1, 1.9],
        fill=DANGER,
    )

    heading(doc, "3. What Was Done")
    para(doc, "The work was not limited to only payment integration. The project was audited, stabilized, redesigned where needed, connected to live backend flows, tested across roles, rebranded, deployed, documented, and pushed to GitHub.")

    heading(doc, "3.1 Payment and Ticket Flow", level=2)
    bullets(
        doc,
        [
            "Integrated Razorpay test payment flow into the passenger ticket booking journey.",
            "Connected payment success to actual ticket generation.",
            "Added/updated required Supabase ticket fields, including amount and ticket status handling.",
            "Verified successful payment creates a real ticket.",
            "Fixed ticket modal scrolling and layout issues.",
            "Improved ticket modal loading behavior and ticket display.",
            "Added clear paid amount, ticket code, validity, expiry, and active status display.",
            "Added two-hour boarding validity rule.",
            "Added expired ticket grouping.",
            "Added used ticket grouping after checker verification.",
            "Added boarding time display after ticket scan/verification.",
            "Added cancellation/refund not allowed messaging.",
        ],
    )

    heading(doc, "3.2 Checker and Verification Flow", level=2)
    bullets(
        doc,
        [
            "Fixed checker dashboard role visibility and role-specific workflow.",
            "Restored QR scanning behavior so tapping scan opens the camera instead of directly showing manual-code errors.",
            "Improved QR scanner camera module and fallback behavior.",
            "Fixed mobile camera behavior after phone testing.",
            "Changed scanner error screens to branded BusLoop UI instead of generic blue/error screens.",
            "Added manual ticket code fallback for cases where camera is blocked or unavailable.",
            "Fixed verification modal and ticket details display.",
            "Added proper valid/denied/rate logic so 0 verified and 0 denied does not show a false 100% rate.",
            "Removed unnecessary route noise from checker page.",
            "Added checker-side show bus QR option so passengers can scan and buy tickets on the bus.",
        ],
    )

    heading(doc, "3.3 Passenger App Improvements", level=2)
    bullets(
        doc,
        [
            "Restored passenger-side bus QR scan use case for buying tickets.",
            "Added live bus test case support.",
            "Improved home page, bus detail, ticket, profile, and review flows.",
            "Fixed QR display issue in ticket modal.",
            "Improved used tickets and reviews flow so rating completed journeys is easier.",
            "Removed redundant review sections and simplified the review UX.",
            "Fixed fare icon/currency display to use rupee context instead of dollar visual language.",
            "Fixed route text truncation and map overflow issues.",
            "Improved map/header visual blending and layout consistency.",
            "Rounded and aligned top/bottom UI surfaces to match the design system.",
            "Cleaned profile section so it looks usable and consistent.",
        ],
    )

    heading(doc, "3.4 Staff App Improvements", level=2)
    bullets(
        doc,
        [
            "Added checker as a first-class staff role on role selection.",
            "Fixed staff role routing after login.",
            "Aligned operator, driver, and checker apps under a shared professional visual language.",
            "Improved operator dashboard layout, analytics, collection totals, and bus management UI.",
            "Fixed mismatch in collection totals and trend numbers.",
            "Fixed Add Bus failure where save gave no result.",
            "Added operator option to view live bus GPS on map.",
            "Improved driver dashboard states and GPS workflow.",
            "Validated driver GPS ping on phone.",
            "Cleaned checker dashboard and no-bus/no-data states.",
            "Reduced unnecessary buttons, duplicate controls, and visual noise.",
        ],
    )

    heading(doc, "3.5 Branding and Deployment", level=2)
    bullets(
        doc,
        [
            "Rebranded the app from BusNOW to BusLoop.",
            "Applied the new BusLoop logo mark across passenger and staff apps.",
            "Updated favicon, app icons, Android launcher icons, and splash assets.",
            "Updated visible app titles and deployed app names.",
            "Renamed Vercel project URLs to BusLoop public URLs.",
            "Removed old public BusNOW URLs from active public routing.",
            "Created and pushed GitHub repository.",
            "Added README with app usage, credentials, and payment test details.",
            "Created client E2E report document.",
        ],
    )

    heading(doc, "4. Before vs After Summary")
    table(
        doc,
        ["Category", "Before", "After"],
        [
            ["Overall State", "Incomplete, inconsistent, not production-ready.", "Stabilized, deployed, documented, and ready for demo/pilot."],
            ["Branding", "BusNOW name and old assets still visible.", "BusLoop branding, logo, favicon, icons, and URLs applied."],
            ["Payment", "Manual/incomplete payment workflow.", "Razorpay test payment working end to end."],
            ["Ticket", "Ticket modal broken/laggy, QR and scroll issues.", "QR ticket generated after payment with validity and expiry."],
            ["Verification", "Checker modal/scanner unreliable.", "QR and manual verification working with branded fallbacks."],
            ["Passenger QR", "Bus QR purchase case missing.", "Passenger can scan bus QR to buy ticket."],
            ["Checker QR", "No clear way to show bus QR to passenger.", "Checker can show bus QR for onboard purchase."],
            ["Operator", "Messy UI and inconsistent analytics.", "Clean dashboard, corrected totals, bus/route/QR/live map flows."],
            ["Driver", "GPS flow needed testing and cleanup.", "Mobile GPS ping tested and working."],
            ["Reviews", "Redundant/confusing review page.", "Past journeys and rating flow simplified."],
            ["Deployment", "Old BusNOW URLs.", "BusLoop user and staff URLs live."],
            ["Documentation", "No complete handover document.", "README and client reports created and pushed."],
        ],
        widths=[1.35, 2.4, 2.5],
    )

    heading(doc, "5. Current Project State")
    callout(
        doc,
        "Current State",
        "The project is now in a client-demo-ready and real-time pilot-ready state. The major passenger and staff journeys are connected, the app is deployed, branding is updated, and end-to-end testing has been documented.",
        fill=SUCCESS,
    )
    table(
        doc,
        ["Current Area", "Status"],
        [
            ["Passenger App", "Live and working at https://busloop-user-app.vercel.app"],
            ["Staff App", "Live and working at https://busloop-staff-app.vercel.app"],
            ["Payment", "Razorpay test mode working with ticket generation."],
            ["Ticket QR", "Generated and displayed after payment."],
            ["Checker Verification", "QR scan and manual ticket verification available."],
            ["Driver GPS", "Phone GPS test confirmed working."],
            ["Operator Live Map", "Available for bus tracking when GPS is active."],
            ["Ticket Validity", "Two-hour boarding window implemented."],
            ["Used/Expired Tickets", "Separated clearly in passenger ticket flow."],
            ["Reviews", "Completed journey rating flow available."],
            ["GitHub", "Project pushed to https://github.com/imadi005/BusLoop"],
            ["Client Documentation", "README and DOCX reports prepared."],
        ],
        widths=[1.8, 4.45],
    )

    heading(doc, "6. Final Tested Login and Payment Reference")
    table(
        doc,
        ["Role", "Email", "Password", "Where to Login"],
        [
            ["Passenger", "passenger@busnow.app", "Passenger@123", "Passenger App"],
            ["Operator", "operator1@busnow.app", "Staff@BusNow2026", "Staff App"],
            ["Driver", "driver1@busnow.app", "Staff@BusNow2026", "Staff App"],
            ["Checker", "checker1@busnow.app", "Staff@BusNow2026", "Staff App"],
        ],
        widths=[1.2, 2.0, 1.8, 1.25],
    )
    table(
        doc,
        ["Payment Field", "Razorpay Test Value"],
        [
            ["Card Number", "5267 3181 8797 5449"],
            ["Expiry", "Any future date, for example 12/30"],
            ["CVV", "Any 3 digits, for example 123"],
            ["Name", "Any name"],
            ["OTP", "123456"],
        ],
        widths=[1.8, 4.45],
    )

    heading(doc, "7. Final Conclusion")
    para(doc, "The project has been moved from an unstable and incomplete state to a functional, branded, tested, and deployed BusLoop product. The core workflows are now understandable for the client, usable by passengers and staff, and ready to demonstrate.")
    callout(
        doc,
        "Final Conclusion",
        "BusLoop is ready for client demo, deployment handover, and real-time pilot use. The remaining production step is to switch Razorpay from test mode to live mode when the client is ready to accept real payments, and to rotate/live-manage final production keys as part of the standard launch process.",
        fill=SUCCESS,
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
