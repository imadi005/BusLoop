from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "BusNOW_E2E_Testing_Report.docx"

BRAND = "EF3E42"
DARK = "101828"
MUTED = "667085"
GREEN = "12B76A"
BLUE = "2563EB"
LIGHT_RED = "FFF1F2"
LIGHT_GREEN = "ECFDF3"
LIGHT_BLUE = "EFF6FF"
TABLE_HEADER = "E8EEF5"
BORDER = "D0D5DD"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=DARK, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def add_table(doc, headers, rows, widths, header_fill=TABLE_HEADER):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True, color=DARK, size=9.5)
        set_cell_shading(hdr[i], header_fill)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            text = str(value)
            set_cell_text(cells[i], text, bold=(i == 0), color=DARK, size=9.2)
            if text.upper() == "PASS" or text.upper() == "WORKING":
                set_cell_shading(cells[i], LIGHT_GREEN)
            elif i == 0:
                set_cell_shading(cells[i], "F9FAFB")
    doc.add_paragraph()
    return table


def add_status_pill_table(doc):
    rows = [
        ("Hosted user app", "WORKING", "Passenger app deployed and used for booking/payment testing."),
        ("Hosted staff app", "WORKING", "Operator, checker, and driver dashboards tested on production URL."),
        ("Razorpay test payment", "WORKING", "Successful payment produced a real BusNOW ticket."),
        ("Supabase live backend", "WORKING", "Ticket, route, staff, bus, and verification records persisted."),
        ("Mobile GPS", "WORKING", "Driver live-location pings confirmed on phone with permission enabled."),
    ]
    add_table(doc, ["Component", "Status", "Verified Result"], rows, [2.0, 1.2, 3.3])


def paragraph(doc, text, style=None, bold_start=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    if bold_start and text.startswith(bold_start):
        r1 = p.add_run(bold_start)
        r1.bold = True
        r2 = p.add_run(text[len(bold_start):])
    else:
        p.add_run(text)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color in [
        ("Heading 1", 16, BRAND),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11.5, DARK),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "BusNOW E2E Testing Report"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = "Prepared for client demo and QA handoff"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(54)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("BusNOW E2E Testing Report")
    run.bold = True
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(BRAND)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run("What was tested and confirmed working across passenger, staff, payment, QR, ticketing, GPS, and admin workflows")
    sr.font.size = Pt(11)
    sr.font.color.rgb = RGBColor.from_string(MUTED)

    meta_rows = [
        ("Date", "06 July 2026"),
        ("Environment", "Hosted Vercel production demo apps"),
        ("User App", "https://busnow-user-app.vercel.app"),
        ("Staff App", "https://busnow-staff-app.vercel.app"),
        ("Backend", "Supabase live project"),
        ("Payment Mode", "Razorpay Test Mode"),
    ]
    add_table(doc, ["Field", "Value"], meta_rows, [1.6, 4.9], header_fill=LIGHT_RED)

    paragraph(
        doc,
        "This report summarizes the final BusNOW end-to-end validation completed for the hosted demo environment. It intentionally focuses on what was tested and confirmed working for client review.",
    )
    doc.add_page_break()

    heading(doc, "1. Overall E2E Status")
    paragraph(
        doc,
        "The main BusNOW user journey is working end to end: a passenger can book a bus ticket, complete Razorpay payment, receive a QR ticket, and have that ticket verified by checker staff.",
    )
    paragraph(
        doc,
        "Staff-side workflows are also working across operator, checker, and driver roles. Operator route, staff, and bus creation were verified after the latest redeploy. Driver GPS was confirmed working on phone with location permission enabled.",
    )
    add_status_pill_table(doc)

    heading(doc, "2. Passenger App Coverage")
    passenger_rows = [
        ("Passenger login", "Passenger account opens the user app successfully.", "PASS"),
        ("Live bus entry", "User can access bus/route booking entry from the app.", "PASS"),
        ("Bus QR booking", "User-side QR scan flow supports bus QR based ticket purchase.", "PASS"),
        ("Route and fare display", "Route, fare, bus, and destination information render correctly.", "PASS"),
        ("Razorpay checkout", "Payment opens and completes in Razorpay Test Mode.", "PASS"),
        ("Payment confirmation", "Successful payment returns to BusNOW and creates the ticket.", "PASS"),
        ("Ticket QR generation", "Ticket modal shows route, status, QR, expiry, and code.", "PASS"),
        ("Ticket validity", "Ticket clearly communicates the two-hour boarding validity.", "PASS"),
        ("Ticket list", "Generated ticket appears in the passenger tickets section.", "PASS"),
        ("Used ticket review", "Used/completed journeys can be rated from tickets/reviews UX.", "PASS"),
        ("Refund policy messaging", "Cancellation/refund not allowed messaging is present.", "PASS"),
    ]
    add_table(doc, ["Tested Area", "Working Behavior", "Status"], passenger_rows, [2.0, 3.7, 0.8])

    heading(doc, "3. Razorpay And Ticketing Coverage")
    payment_rows = [
        ("Razorpay order/payment", "Domestic test card payment completes successfully.", "PASS"),
        ("Payment verification", "Payment success is verified through the backend flow.", "PASS"),
        ("Ticket persistence", "Ticket record is saved in Supabase with payment reference.", "PASS"),
        ("Amount handling", "Ticket amount is saved and shown correctly as INR fare.", "PASS"),
        ("QR payload", "Generated QR contains the ticket code used for validation.", "PASS"),
        ("Post-payment UX", "Success ticket modal loads after payment and remains scrollable.", "PASS"),
    ]
    add_table(doc, ["Tested Area", "Working Behavior", "Status"], payment_rows, [2.0, 3.7, 0.8])
    doc.add_page_break()

    heading(doc, "4. Checker App Coverage")
    checker_rows = [
        ("Checker login", "Checker role opens the correct staff dashboard.", "PASS"),
        ("Assigned bus context", "Checker sees assigned/active bus context for verification.", "PASS"),
        ("Camera QR scanner", "Camera scanner works on supported mobile browser setup.", "PASS"),
        ("Manual code entry", "Manual ticket-code verification works as fallback.", "PASS"),
        ("Valid ticket lookup", "Valid ticket details load with passenger and route data.", "PASS"),
        ("Mark verified", "Checker can mark an active ticket as boarded/used.", "PASS"),
        ("Boarding time", "Passenger-side ticket shows boarding/scan time after verification.", "PASS"),
        ("Duplicate prevention", "Already-used tickets cannot be marked again.", "PASS"),
        ("Invalid/expired handling", "Invalid or expired tickets are rejected with user feedback.", "PASS"),
        ("Show bus QR", "Checker can show bus QR so passengers can buy a ticket onboard.", "PASS"),
    ]
    add_table(doc, ["Tested Area", "Working Behavior", "Status"], checker_rows, [2.0, 3.7, 0.8])

    heading(doc, "5. Operator App Coverage")
    operator_rows = [
        ("Operator login", "Operator account opens the correct dashboard.", "PASS"),
        ("Analytics dashboard", "Collections, trips, tickets, buses, and rating KPIs render.", "PASS"),
        ("Collection totals", "Daily collection and trend values use ticket amount data.", "PASS"),
        ("Route creation", "New route can be created and appears in route list.", "PASS"),
        ("Staff creation", "New driver/checker staff account can be created.", "PASS"),
        ("Bus creation", "New bus can be created after redeploy and appears in fleet.", "PASS"),
        ("Bus QR generation", "New bus opens a QR modal with downloadable/copyable payload.", "PASS"),
        ("Route assignment", "Operator can assign route from the fleet controls.", "PASS"),
        ("Live bus map", "Operator can view live GPS map for active buses.", "PASS"),
        ("Inactive live guard", "Live button remains disabled for buses without active trip.", "PASS"),
        ("Settings", "Operator settings surface is accessible and aligned with staff UI.", "PASS"),
    ]
    add_table(doc, ["Tested Area", "Working Behavior", "Status"], operator_rows, [2.0, 3.7, 0.8])
    doc.add_page_break()

    heading(doc, "6. Driver App And GPS Coverage")
    driver_rows = [
        ("Driver login", "Driver account opens the driver dashboard.", "PASS"),
        ("Assigned bus", "Driver can see assigned/available bus information.", "PASS"),
        ("Active trip", "Active trip details display with route and bus information.", "PASS"),
        ("GPS permission", "Location permission flow works on phone browser.", "PASS"),
        ("Live pings", "GPS pings are sent from phone and bus becomes live.", "PASS"),
        ("Passenger visibility", "Live bus state supports passenger map/booking flow.", "PASS"),
        ("Operator visibility", "Active bus can be viewed from operator live map.", "PASS"),
    ]
    add_table(doc, ["Tested Area", "Working Behavior", "Status"], driver_rows, [2.0, 3.7, 0.8])

    heading(doc, "7. Role And Access Coverage")
    access_rows = [
        ("Passenger role", "Passenger sees passenger tabs and ticket/review/profile flows.", "PASS"),
        ("Checker role", "Checker sees scan-first verification workflow.", "PASS"),
        ("Driver role", "Driver sees trip and location-sharing workflow.", "PASS"),
        ("Operator role", "Operator sees fleet, analytics, routes, staff, and settings.", "PASS"),
        ("Staff separation", "Staff roles are separated from passenger app experience.", "PASS"),
        ("Role display", "Profile/dashboard role labels now align with logged-in role.", "PASS"),
    ]
    add_table(doc, ["Tested Area", "Working Behavior", "Status"], access_rows, [2.0, 3.7, 0.8])

    heading(doc, "8. Data Created During E2E")
    data_rows = [
        ("Route", "E2E Origin 050151 to E2E Destination 050151", "Created"),
        ("Staff", "E2E Driver 050151", "Created"),
        ("Bus", "E2E Fixed Bus 039156", "Created"),
        ("Ticket", "045163e5-2343-487b-a6ef-02915a59ce97", "Generated and verified"),
    ]
    add_table(doc, ["Record Type", "Record", "Result"], data_rows, [1.5, 3.7, 1.3])

    heading(doc, "9. Final Working Summary")
    summary_rows = [
        ("Passenger booking", "Working from route/bus selection through payment and ticket display."),
        ("Payment integration", "Razorpay Test Mode payment and backend verification are working."),
        ("Ticket lifecycle", "Active, used, validity, expiry, QR, and duplicate-use prevention are working."),
        ("Checker verification", "QR/manual verification and onboard purchase QR support are working."),
        ("Operator management", "Analytics, route creation, staff creation, bus creation, QR generation, and live-map controls are working."),
        ("Driver GPS", "Phone-based GPS pings and live bus tracking are working."),
        ("Client demo readiness", "The core product workflow is ready for client demonstration in the hosted environment."),
    ]
    add_table(doc, ["Workflow", "Confirmed Working Result"], summary_rows, [2.1, 4.4])

    paragraph(
        doc,
        "Security-sensitive credentials are intentionally excluded from this client-facing report.",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
