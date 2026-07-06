from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "BusLoop_Client_Project_Report.docx"
LOGO = ROOT / "busnow-user-app" / "public" / "busloop-mark.png"

NAVY = RGBColor(11, 37, 69)
RED = RGBColor(239, 61, 65)
GREEN = RGBColor(18, 183, 106)
BLUE = RGBColor(37, 99, 235)
MUTED = RGBColor(98, 111, 134)
LIGHT = "F5F7FB"
HEADER = "E8EEF5"
SUCCESS = "EAF8F1"
WARN = "FFF7E6"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None, header_fill=HEADER):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], header_fill)
        set_cell_text(hdr[i], h, bold=True, color=NAVY)
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, item in enumerate(row):
            set_cell_text(cells[i], item)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if widths:
        set_table_widths(table, widths)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        run.font.color.rgb = NAVY if level == 1 else BLUE
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.color.rgb = NAVY
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_callout(doc, title, body, fill=SUCCESS):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = NAVY
    r.font.size = Pt(11)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(body)
    set_table_widths(table, [6.25])
    doc.add_paragraph()


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Heading 1", 16, NAVY),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11.5, NAVY),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True


def set_header_footer(doc):
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("BusLoop Client Report")
    r.font.size = Pt(8.5)
    r.font.color.rgb = MUTED
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = fp.add_run("BusLoop | Passenger Ticketing, Staff Operations, Live Tracking")
    rf.font.size = Pt(8.5)
    rf.font.color.rgb = MUTED


def build():
    doc = Document()
    setup_styles(doc)
    set_header_footer(doc)

    # Cover
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(1.45))
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("BusLoop Client Project Report")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = NAVY
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Live demo access, workflows, test credentials, payment testing, and final E2E validation")
    sr.font.size = Pt(11.5)
    sr.font.color.rgb = MUTED
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("Prepared for client review | July 6, 2026")
    mr.font.size = Pt(10)
    mr.font.color.rgb = MUTED

    add_callout(
        doc,
        "Final Status: Ready for deployment and real-time use",
        "The BusLoop passenger and staff applications have been deployed, rebranded, tested across major user journeys, and validated with live Supabase data, Razorpay test payment, QR ticket generation, checker verification, staff workflows, and mobile camera/GPS behavior.",
        fill=SUCCESS,
    )

    add_table(
        doc,
        ["Project Link", "Details"],
        [
            ["Passenger App", "https://busloop-user-app.vercel.app"],
            ["Staff App", "https://busloop-staff-app.vercel.app"],
            ["GitHub Repository", "https://github.com/imadi005/BusLoop"],
            ["Backend", "Supabase Auth, Database, Realtime, and Edge Functions"],
            ["Payments", "Razorpay test-mode payment integration"],
        ],
        widths=[1.75, 4.5],
    )

    doc.add_page_break()

    add_heading(doc, "1. Executive Summary")
    add_para(
        doc,
        "BusLoop is a complete bus ticketing and live tracking system. It has a passenger app for ticket booking and a staff app for bus operators, drivers, and checkers. The system is designed so passengers can find buses, buy QR tickets online, board within the valid time window, and get their ticket verified by staff."
    )
    add_para(
        doc,
        "On the staff side, operators can manage the service, drivers can start trips and share live GPS location, and checkers can validate passenger tickets using QR scanning or manual ticket code entry."
    )
    add_bullets(
        doc,
        [
            "Passenger ticket booking and Razorpay payment flow are working.",
            "Ticket generation with QR code is working after successful payment.",
            "Checker-side QR verification and manual ticket verification are working.",
            "Driver live GPS flow is working on mobile.",
            "Operator dashboard, buses, routes, QR, repair, analytics, and live map flows are available.",
            "Ticket validity, expiry, used tickets, boarding time, no-refund messaging, and review flows are implemented.",
            "The product has been rebranded from BusNOW to BusLoop with updated logo, favicon, app titles, and deployed URLs.",
        ],
    )

    add_heading(doc, "2. Application Access")
    add_table(
        doc,
        ["App", "Who Uses It", "Live URL"],
        [
            ["Passenger App", "Passengers / riders", "https://busloop-user-app.vercel.app"],
            ["Staff App", "Operator, Driver, Checker", "https://busloop-staff-app.vercel.app"],
        ],
        widths=[1.35, 2.0, 2.9],
    )

    add_heading(doc, "3. Tested Demo Login Credentials")
    add_table(
        doc,
        ["Role", "Email", "Password", "App"],
        [
            ["Passenger", "passenger@busnow.app", "Passenger@123", "Passenger App"],
            ["Operator", "operator1@busnow.app", "Staff@BusNow2026", "Staff App"],
            ["Driver", "driver1@busnow.app", "Staff@BusNow2026", "Staff App"],
            ["Checker", "checker1@busnow.app", "Staff@BusNow2026", "Staff App"],
        ],
        widths=[1.25, 2.0, 1.8, 1.2],
    )

    add_heading(doc, "4. Razorpay Test Payment Details")
    add_para(doc, "The payment flow was tested using Razorpay test mode. The following details can be used for client demo payment testing.")
    add_table(
        doc,
        ["Payment Field", "Test Value"],
        [
            ["Card Number", "5267 3181 8797 5449"],
            ["Expiry", "Any future date, for example 12/30"],
            ["CVV", "Any 3 digits, for example 123"],
            ["Name", "Any name"],
            ["OTP", "123456"],
        ],
        widths=[1.8, 4.45],
    )

    add_heading(doc, "5. Passenger App Workflow")
    add_para(doc, "The passenger app is built for the rider journey from discovery to payment, ticket generation, boarding, and review.")
    add_steps(
        doc,
        [
            "Passenger opens the Passenger App URL.",
            "Passenger logs in with the demo passenger account.",
            "Passenger views available buses and live bus status.",
            "Passenger can scan a bus QR code to directly open that bus booking page.",
            "Passenger selects the bus and books a ticket.",
            "Razorpay payment modal opens and payment is completed using test card details.",
            "After successful payment, the app creates a paid ticket and displays the QR ticket modal.",
            "The QR ticket can be saved and shown to the checker while boarding.",
            "Tickets are shown in active, used, and expired sections.",
            "After a trip is verified/used, the passenger can rate the trip.",
        ],
    )

    add_heading(doc, "Passenger Features Confirmed")
    add_bullets(
        doc,
        [
            "Login and passenger session flow.",
            "Bus list and bus detail page.",
            "Bus QR scan path for buying tickets.",
            "Razorpay test payment.",
            "Ticket generation after payment.",
            "QR ticket rendering and save option.",
            "Two-hour ticket validity notice.",
            "Expired ticket grouping.",
            "Used ticket grouping.",
            "Boarding time shown after checker verification.",
            "Cancellation/refund not allowed message.",
            "Trip rating from used tickets and reviews page.",
            "Profile section UI cleanup and usable settings layout.",
        ],
    )

    add_heading(doc, "6. Checker Workflow")
    add_para(doc, "The checker workflow is designed for staff who verify passenger tickets during boarding.")
    add_steps(
        doc,
        [
            "Checker opens Staff App and selects Checker role.",
            "Checker logs in with checker credentials.",
            "Checker sees assigned bus and verification statistics.",
            "Checker can scan passenger QR ticket using the camera.",
            "If camera is unavailable, checker can manually enter the ticket code.",
            "Valid ticket details are shown before marking the ticket as verified.",
            "Checker marks ticket verified.",
            "Ticket status changes to used and boarding time is recorded.",
            "Checker can open the current bus QR code so passengers without tickets can scan and buy.",
        ],
    )
    add_table(
        doc,
        ["Checker Case", "Result"],
        [
            ["Valid active ticket", "Ticket details open correctly and can be verified."],
            ["Already used ticket", "Ticket is no longer treated as a fresh active ticket."],
            ["Expired ticket", "Ticket is separated from active valid tickets."],
            ["No code entered", "User receives a proper validation message."],
            ["Camera blocked/unavailable", "App shows branded fallback with manual entry option."],
            ["Mobile browser camera", "Camera module was fixed and tested on phone."],
            ["No assigned/active bus", "UI avoids unnecessary route noise and keeps state clear."],
        ],
        widths=[2.1, 4.15],
    )

    add_heading(doc, "7. Driver Workflow")
    add_para(doc, "The driver workflow supports real-time bus visibility by sending GPS updates during an active trip.")
    add_steps(
        doc,
        [
            "Driver opens Staff App and selects Driver role.",
            "Driver logs in with driver credentials.",
            "Driver selects assigned bus.",
            "Driver starts a trip.",
            "Driver grants location permission on mobile.",
            "Live GPS updates are sent while the trip is active.",
            "Passenger and operator views can use this live location information.",
            "Driver ends the trip when the journey is complete.",
        ],
    )
    add_table(
        doc,
        ["Driver Case", "Result"],
        [
            ["Mobile GPS permission", "Tested on phone and confirmed working."],
            ["Start trip", "Driver can start live trip."],
            ["Resume active trip", "Driver dashboard resumes existing active trip state."],
            ["End trip", "Trip can be completed."],
            ["No active bus/state", "UI shows clean states without unnecessary controls."],
        ],
        widths=[2.1, 4.15],
    )

    add_heading(doc, "8. Operator Workflow")
    add_para(doc, "The operator workflow is for fleet and route management, live monitoring, staff overview, and business analytics.")
    add_steps(
        doc,
        [
            "Operator opens Staff App and selects Operator role.",
            "Operator logs in with operator credentials.",
            "Operator sees dashboard analytics and collection data.",
            "Operator manages buses, routes, and staff sections.",
            "Operator can add buses and routes.",
            "Operator can open bus QR codes for passenger booking.",
            "Operator can view bus live GPS location on map when the driver is sharing location.",
            "Operator can inspect bus status, active/delayed/inactive state, and repair actions.",
        ],
    )
    add_table(
        doc,
        ["Operator Case", "Result"],
        [
            ["Analytics dashboard", "Collection and ticket counts were corrected for consistency."],
            ["Add bus flow", "Save issue was fixed; button state, save behavior, and feedback were corrected."],
            ["Bus QR", "Bus QR generation/view option available."],
            ["Live bus map", "Operator can open live GPS map for bus tracking."],
            ["Routes", "Route creation and management available."],
            ["Staff overview", "Operator can view operational staff data."],
            ["UI consistency", "Operator screens were cleaned up and aligned with the shared BusLoop style."],
        ],
        widths=[2.1, 4.15],
    )

    add_heading(doc, "9. Ticket and Payment Rules")
    add_table(
        doc,
        ["Rule", "Behavior"],
        [
            ["Payment required", "Ticket is generated only after successful payment."],
            ["Ticket validity", "Passenger must board within 2 hours of booking."],
            ["Expired ticket", "Expired tickets are moved away from active tickets."],
            ["Used ticket", "Ticket becomes used after checker verification."],
            ["Boarding time", "The time of checker verification is stored and shown to the passenger."],
            ["Refund/cancellation", "Cancellation and refund are not allowed in the current flow."],
            ["QR ticket", "Ticket QR is created and displayed after successful booking."],
        ],
        widths=[1.8, 4.45],
    )

    add_heading(doc, "10. End-to-End Test Results")
    add_para(doc, "The following results summarize the final end-to-end checks performed across passenger, payment, checker, driver, and operator workflows.")
    add_table(
        doc,
        ["Area Tested", "Status", "Result Summary"],
        [
            ["Passenger login", "PASS", "Passenger account can log in and access the user app."],
            ["Staff login", "PASS", "Operator, Driver, and Checker accounts can log in through staff app."],
            ["Role routing", "PASS", "Staff users are routed to the correct role dashboard."],
            ["Razorpay payment", "PASS", "Test card payment completed successfully."],
            ["Ticket creation", "PASS", "Paid ticket generated after payment."],
            ["Ticket QR display", "PASS", "QR ticket modal renders and ticket details are visible."],
            ["Ticket save option", "PASS", "Save QR to device option is available."],
            ["Checker scan/manual verify", "PASS", "Ticket can be checked and marked verified."],
            ["Used ticket state", "PASS", "Verified tickets move to used state."],
            ["Boarding time", "PASS", "Boarding time is shown after verification."],
            ["Ticket expiry", "PASS", "Two-hour validity and expired ticket grouping implemented."],
            ["Cancellation/refund rule", "PASS", "No cancellation/refund rule is shown to users."],
            ["Passenger bus QR scan", "PASS", "Bus QR scan flow restored for ticket purchase."],
            ["Checker show bus QR", "PASS", "Checker can show bus QR for passengers who have not bought a ticket."],
            ["Camera module", "PASS", "Camera unavailable/errors now use branded UI and manual fallback."],
            ["Mobile camera", "PASS", "Camera behavior was fixed and validated on phone."],
            ["Driver GPS", "PASS", "GPS ping was checked on phone and confirmed working."],
            ["Operator add bus", "PASS", "Add bus failure was fixed and redeployed."],
            ["Operator live map", "PASS", "Operator can view live bus location when GPS is available."],
            ["Review flow", "PASS", "Used/past journeys can be rated with simplified review UX."],
            ["UI consistency", "PASS", "Passenger and staff apps were aligned under a cleaner BusLoop visual style."],
            ["Branding", "PASS", "BusLoop name, logo, favicon, and URLs are applied."],
            ["Deployment URLs", "PASS", "New BusLoop URLs are live; old public BusNOW URLs return 404."],
        ],
        widths=[1.75, 0.75, 3.75],
    )

    add_heading(doc, "11. Security and Production Handling")
    add_para(doc, "The app was prepared with safer production handling for credentials and privileged operations.")
    add_bullets(
        doc,
        [
            "Razorpay secret key is handled server-side through Supabase Edge Function secrets.",
            "Supabase service-role key is handled server-side and not exposed in frontend code.",
            "Passenger app uses the public Supabase anon key only.",
            "Staff role access is separated between operator, driver, and checker dashboards.",
            "Ticket creation and payment verification are routed through backend functions.",
            "Environment files, Vercel metadata, build folders, node_modules, and local Supabase temp files are excluded from Git.",
        ],
    )

    add_heading(doc, "12. Final Client Conclusion")
    add_callout(
        doc,
        "Deployment Readiness Conclusion",
        "BusLoop is ready for deployment and real-time use. The passenger app, staff app, payment flow, ticket generation, QR verification, live GPS tracking, staff workflows, and role-based user journeys have been tested and validated. The system is suitable for client demo, operational pilot, and real-world rollout with the configured Supabase and Razorpay environments.",
        fill=SUCCESS,
    )
    add_para(
        doc,
        "Recommended next operational step: before moving from test mode to live payments, replace Razorpay test keys with live keys, keep secrets only in Supabase Edge Function secrets, and confirm final domain ownership if a custom business domain is required."
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
